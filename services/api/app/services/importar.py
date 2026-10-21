"""Leer destinatarios desde un Excel o un CSV.

Lo que se acepta es lo que la gente realmente tiene: la exportación de un sistema con
encabezados ("Cliente", "Celular", ...), o una lista pelada de teléfono y nombre sin
encabezado. Solo se devuelve lo que hay en el archivo; la normalización del teléfono y
los duplicados los resuelve quien carga, igual que con la lista pegada a mano.
"""

from __future__ import annotations

import csv
import io
import unicodedata
import zipfile

MAX_FILAS = 5000

_TELEFONO = ("telefono", "tel", "celular", "cel", "movil", "phone", "numero", "nro", "whatsapp", "contacto tel")
_NOMBRE = ("nombre", "cliente", "name", "contacto", "razon social", "titular")


class ArchivoInvalido(ValueError):
    """El motivo se muestra tal cual a quien subió el archivo."""


def _sin_tildes(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn").lower().strip()


def _celda(valor: object) -> str:
    """Excel guarda los teléfonos como número (981123456.0): se devuelven como texto sin el .0."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _filas_csv(contenido: bytes) -> list[list[str]]:
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    return [[_celda(c) for c in fila] for fila in csv.reader(io.StringIO(texto), dialecto)]


MAX_DESCOMPRIMIDO = 60 * 1024 * 1024
MAX_ENTRADAS = 300


def _verificar_zip(contenido: bytes) -> None:
    """Un .xlsx es un zip: uno de 5 MB puede esconder gigas (zip bomb) que se expanden en memoria."""
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            infos = z.infolist()
    except zipfile.BadZipFile as exc:
        raise ArchivoInvalido("No se pudo abrir el Excel. ¿Es un .xlsx?") from exc
    if len(infos) > MAX_ENTRADAS or sum(i.file_size for i in infos) > MAX_DESCOMPRIMIDO:
        raise ArchivoInvalido("El Excel es demasiado grande por dentro. Guardalo de nuevo o usá un CSV.")


def _filas_excel(contenido: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    _verificar_zip(contenido)
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl lanza de todo con un archivo roto
        raise ArchivoInvalido("No se pudo abrir el Excel. ¿Es un .xlsx?") from exc
    hoja = libro.worksheets[0]
    filas = []
    for fila in hoja.iter_rows(values_only=True):
        filas.append([_celda(c) for c in fila])
        if len(filas) > MAX_FILAS + 1:
            break
    libro.close()
    return filas


def _columna(encabezado: list[str], claves: tuple[str, ...]) -> int | None:
    normal = [_sin_tildes(h) for h in encabezado]
    for clave in claves:  # en orden de preferencia: "telefono" gana sobre "contacto"
        for i, h in enumerate(normal):
            if h == clave or h.startswith(clave + " ") or h.startswith(clave + "_") or h.startswith(clave + "."):
                return i
    for clave in claves:
        for i, h in enumerate(normal):
            if clave in h:
                return i
    return None


def leer_filas(contenido: bytes, nombre_archivo: str) -> list[tuple[str, str | None]]:
    """[(teléfono tal cual, nombre o None)]. Lanza ArchivoInvalido con un motivo legible."""
    return [(tel, nom) for tel, nom, _ in leer_filas_con_datos(contenido, nombre_archivo)]


def leer_filas_con_datos(contenido: bytes, nombre_archivo: str) -> list[tuple[str, str | None, dict[str, str]]]:
    """[(teléfono, nombre o None, datos extra)]: las demás columnas («Modelo», «Sucursal»…) son datos para el guion
    (ver app/datos_cliente.py). Sin encabezado reconocible no hay datos extra."""
    if not contenido:
        raise ArchivoInvalido("El archivo está vacío.")
    nombre = (nombre_archivo or "").lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        filas = _filas_excel(contenido)
    elif nombre.endswith((".csv", ".txt")):
        filas = _filas_csv(contenido)
    elif nombre.endswith(".xls"):
        raise ArchivoInvalido("Es un Excel viejo (.xls). Guardalo como .xlsx o como CSV y volvé a subirlo.")
    else:
        raise ArchivoInvalido("Subí un .xlsx o un .csv.")

    filas = [f for f in filas if any(c for c in f)]
    if not filas:
        raise ArchivoInvalido("El archivo no tiene filas.")

    col_tel = _columna(filas[0], _TELEFONO)
    col_nom = _columna(filas[0], _NOMBRE)
    encabezado: list[str] = []
    if col_tel is not None:
        encabezado, cuerpo = filas[0], filas[1:]
    else:
        # Sin encabezado reconocible: teléfono en la primera columna y nombre en la segunda.
        col_tel, col_nom, cuerpo = 0, 1, filas
    if col_nom == col_tel:
        col_nom = None

    if len(cuerpo) > MAX_FILAS:
        raise ArchivoInvalido(f"El archivo tiene más de {MAX_FILAS} filas. Partilo en varios.")

    from app.datos_cliente import datos_de_fila

    ignorar = {i for i in (col_tel, col_nom) if i is not None}
    salida: list[tuple[str, str | None, dict[str, str]]] = []
    for fila in cuerpo:
        tel = fila[col_tel] if col_tel < len(fila) else ""
        nom = fila[col_nom] if col_nom is not None and col_nom < len(fila) else ""
        if tel or nom:
            salida.append((tel, nom or None, datos_de_fila(encabezado, fila, ignorar) if encabezado else {}))
    return salida
